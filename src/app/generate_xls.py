from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.cell.text import InlineFont
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.styles import (Alignment,
                             Border,
                             Color,
                             Side, 
                             NamedStyle,
                             Font,
                             PatternFill)
from openpyxl.worksheet.worksheet import Worksheet


from app.settings import APP_DIR, STATIC_DIR

normalize: NamedStyle = NamedStyle(
    name="normalize",
)
normalize.font = Font(name="Arial", size=14, color="000000")
normalize.fill = PatternFill(fill_type=None, start_color="FFFFFF")
normalize.border = Border()
normalize.alignment = Alignment(horizontal='center', vertical='center')

tb_border: NamedStyle = NamedStyle(
    name="tb_border",
)
tb_border.border = Border(top=Side("thin", "000000"), bottom=Side("thin", "000000"))

def adjust_col_width(ws):
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2) * 1.5
        ws.column_dimensions[column_letter].width = adjusted_width

def apply_style(items: tuple, style_name: str):
    for item in items:
        for cell in item:
            cell.style = style_name

data = {
    "system": "НКУ",
    "input": {
        "quantity": 2,
        "control_signals": True,
        "cs_protocol": "Modbus RTU",
        "io_modul_type": "ЭНМВ-1",
        "mesurment": True,
        "mesurment_protocol": "Modbus RTU",
        "mip_type": "ЭНИП-2"
    },
    "output": {
        "quantity": 5,
        "control_signals": False,
        "io_modul_type": "ЭНМВ-1",
        "mesurment": True,
        "mip": "ЭНИП-2"
    }
}
CONTROL_COMANDS = ["Команда включить", "Команда отключить"]
INPUT_SIGNALS = [
    'Положение выключателя "включен"',
    'Положение выключателя "выключен"',
    "Аварийное отключение",
]
MESURMENT_SIGNALS = [
    "Ток - фаза A",
    "Ток - фаза B",
    "Ток - фаза C",
    
]

def format(ws:Worksheet):
    ws.insert_rows(1, 2)
    ws.merge_cells("A1:G1")
    ws.row_dimensions[1].height = 117
    ws["A1"].value = CellRichText(
        "ОБЩЕСТВО С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ\n",
        "ПРОИЗВОДСТВЕННОЕ ОБЪЕДИНЕНИЕ\n",
        TextBlock(InlineFont(b=True), "«ВЫСОКОВОЛЬТНЫЕ ЭЛЕКТРОТЕХНИЧЕСКИЕ АППАРАТЫ»")
    )
    ws["A2"].alignment = Alignment(horizontal="right", vertical="bottom")
    logo = Image(STATIC_DIR/ "img/logo.png")
    logo.width = 253
    logo.height = 60
    ws.add_image(logo, "A1")
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 95
    apply_style(ws["A2:HG2"], "tb_border")
    ws["A2"].value = "Перечень сигналов"
    ws["A2"].font = Font(name="Arial", size=14, b=True, color="000000")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    

    


def generate(data: dict):
    wb = Workbook()
    wb.add_named_style(normalize)
    wb.add_named_style(tb_border)
    ws = wb.active
    ws.append(
        [
            "№ п/п",
            "Система",
            "Наименование", 
            "Наименование сигнала",
            "Источник сигнала",
            "Тип сигнала",
            "Примечание"
        ]
    )
    row_counter = 0
    for i in range(1, data["input"]["quantity"]+1):
        if data["input"]["control_signals"]:
            for comand in CONTROL_COMANDS:
                row_counter += 1
                ws.append(
                    [
                        row_counter,
                        data["system"],
                        f"Ввод {i}",
                        comand,
                        data["input"]["io_modul_type"],
                        "ТУ",
                        ""
                    ]
                )
        for signal in INPUT_SIGNALS:
            row_counter += 1
            ws.append(
                [
                    row_counter,
                    data["system"],
                    f"Ввод {i}",
                    signal,
                    data["input"]["io_modul_type"],
                    "ТС",
                    ""
                ]
            )
        if data["input"]["mesurment"]:
            for signal in MESURMENT_SIGNALS:
                row_counter += 1
                ws.append(
                    [
                        row_counter,
                        data["system"],
                        f"Ввод {i}",
                        signal,
                        data["input"]["mip_type"],
                        "ТИ",
                        ""
                    ]
                )
            
    for i in range(data["output"]["quantity"]):
        if data["output"]["control_signals"]:
            for comand in CONTROL_COMANDS:
                row_counter += 1
                ws.append(
                    [
                        row_counter,
                        data["system"],
                        f"Ввод {i}",
                        comand,
                        data["input"]["io_modul_type"],
                        "ТУ",
                        ""
                    ]
                )
        for signal in INPUT_SIGNALS:
            row_counter += 1
            ws.append(
                [
                    row_counter,
                    data["system"],
                    f"Ввод {i}",
                    signal,
                    data["input"]["io_modul_type"],
                    "ТС",
                    ""
                ]
            )
        if data["input"]["mesurment"]:
            for signal in MESURMENT_SIGNALS:
                row_counter += 1
                ws.append(
                    [
                        row_counter,
                        data["system"],
                        f"Ввод {i}",
                        signal,
                        data["input"]["mip_type"],
                        "ТИ",
                        ""
                    ]
                )
    apply_style(ws.rows, "normalize")
    adjust_col_width(ws)
    format(ws)
    wb.save(APP_DIR.parent.parent / "temp/example.xlsx")


if __name__ == "__main__":
    generate(data)
    