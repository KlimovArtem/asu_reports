from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.cell.text import InlineFont
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, Field

from app.domain import interfaces
from app.settings import APP_DIR, STATIC_DIR
from app.utils.tkinter_utils import adjust_col_width, apply_style, normalize, tb_border


class SignalsListSubdata(BaseModel):
    quantity: int
    di_module_type: str
    control: bool
    do_module_type: str
    measurements: bool
    ai_module_type: str


class SignalsListData(BaseModel):
    system: str
    supplys: SignalsListSubdata
    feeders: SignalsListSubdata

CONTROL_COMANDS = ["Команда включить", "Команда отключить"]
INPUT_SIGNALS = [
    'Положение выключателя "включен"',
    'Положение выключателя "выключен"',
    "Аварийное отключение",
]
MESURMENT_SIGNALS = [
    "Ток - фаза A",
    "Ток - фаза B",
    "Ток - фаза C",
    
]


class BaseReport(interfaces.ReportInterface):

    def __init__(self, title: str = ""):
        self.title: str = title
        self.content: Any = None
        self.data_scheme: BaseModel | None = None
        
        

class XLSReport(BaseReport):
    def __init__(self, title: str):
        super().__init__(title=title)
        self.content: Workbook = Workbook()
    
    def format(self):
        pass
    
    def generate(self, data):
        pass

    def save(self, path):
        pass



class SignalsList(XLSReport):
    def __init__(self, title: str):
        super().__init__(title=title)
        self.data_scheme = SignalsListData

    def format(self, ws:Worksheet) -> None:
        ws.insert_rows(1, 2)
        ws.merge_cells("A1:G1")
        ws.row_dimensions[1].height = 117
        ws["A1"].value = CellRichText(
            "ОБЩЕСТВО С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ\n",
            "ПРОИЗВОДСТВЕННОЕ ОБЪЕДИНЕНИЕ\n",
            TextBlock(InlineFont(b=True), "«ВЫСОКОВОЛЬТНЫЕ ЭЛЕКТРОТЕХНИЧЕСКИЕ АППАРАТЫ»")
        )
        ws["A2"].alignment = Alignment(horizontal="right", vertical="bottom")
        logo = Image(STATIC_DIR/ "img/logo.png")
        logo.width = 253
        logo.height = 60
        ws.add_image(logo, "A1")
        ws.merge_cells("A2:G2")
        ws.row_dimensions[2].height = 95
        apply_style(ws["A2:HG2"], "tb_border")
        ws["A2"].value = "Перечень сигналов"
        ws["A2"].font = Font(name="Arial", size=14, b=True, color="000000")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    def generate(self, data: SignalsListData) -> None:
        self.content.add_named_style(normalize)
        self.content.add_named_style(tb_border)
        ws = self.content.active
        ws.append(
            [
                "№ п/п",
                "Система",
                "Наименование", 
                "Наименование сигнала",
                "Источник сигнала",
                "Тип сигнала",
                "Примечание"
            ]
        )
        row_counter = 0
        for i in range(1, data.supplys.quantity + 1):
            if data.supplys.control:
                for comand in CONTROL_COMANDS:
                    row_counter += 1
                    ws.append(
                        [
                            row_counter,
                            data.system,
                            f"Ввод {i}",
                            comand,
                            data.supplys.do_module_type,
                            "ТУ",
                            ""
                        ]
                    )
            for signal in INPUT_SIGNALS:
                row_counter += 1
                ws.append(
                    [
                        row_counter,
                        data.system,
                        f"Ввод {i}",
                        signal,
                        data.supplys.di_module_type,
                        "ТС",
                        ""
                    ]
                )
            if data.supplys.measurements:
                for signal in MESURMENT_SIGNALS:
                    row_counter += 1
                    ws.append(
                        [
                            row_counter,
                            data.system,
                            f"Ввод {i}",
                            signal,
                            data.supplys.ai_module_type,
                            "ТИ",
                            ""
                        ]
                    )
                
        for i in range(1, data.feeders.quantity + 1 ):
            if data.feeders.control:
                for comand in CONTROL_COMANDS:
                    row_counter += 1
                    ws.append(
                        [
                            row_counter,
                            data.system,
                            f"ОЛ {i}",
                            comand,
                            data.feeders.do_module_type,
                            "ТУ",
                            ""
                        ]
                    )
            for signal in INPUT_SIGNALS:
                row_counter += 1
                ws.append(
                    [
                        row_counter,
                        data.system,
                        f"ОЛ {i}",
                        signal,
                        data.feeders.di_module_type,
                        "ТС",
                        ""
                    ]
                )
            if data.feeders.measurements:
                for signal in MESURMENT_SIGNALS:
                    row_counter += 1
                    ws.append(
                        [
                            row_counter,
                            data.system,
                            f"ОЛ {i}",
                            signal,
                            data.feeders.ai_module_type,
                            "ТИ",
                            ""
                        ]
                    )
        apply_style(ws.rows, "normalize")
        adjust_col_width(ws)
        self.format(ws)

    def save(self, path: Path | str) -> None:
        self.content.save(path / f"{self.title}.xlsx")


if __name__ == "__main__":
    signals_list_example = SignalsList(title="Пример отчёта")
    signals_list_example.save(APP_DIR.parent.parent / "temp")